#!/usr/bin/env python3
"""
Runs MobileVideoGPT 0.5B (with LoRA adapters) inference across multiple exercises
at different temperatures and produces an Excel report with per-video metrics.

Usage:
    python experiments/vary_inference.py
    python experiments/vary_inference.py --videos_per_exercise 10 --device cpu
    python experiments/vary_inference.py --dataset_path dataset/QEVD-14-CLEANED --output results/temp_sweep.xlsx
"""

import sys
import os
import warnings
import logging
import argparse
import json
import re
import time

os.environ["PYTHONWARNINGS"] = "ignore"
warnings.filterwarnings("ignore")
logging.getLogger("mmengine").setLevel(logging.CRITICAL)
logging.getLogger("transformers").setLevel(logging.CRITICAL)

import torch
import numpy as np
from pathlib import Path
from tqdm import tqdm
from typing import List, Dict, Tuple, Optional

from transformers import AutoTokenizer, AutoModelForCausalLM, AutoConfig
from peft import PeftModel
from huggingface_hub import hf_hub_download

# Add workspace root to path
sys.path.insert(0, str(Path(__file__).parent.parent))

from mobilevideogpt.utils import preprocess_input

# Defaults
DEFAULT_DATASET = "dataset/QEVD-14-CLEANED"
DEFAULT_BASE_MODEL = "Amshaker/Mobile-VideoGPT-0.5B"
DEFAULT_LORA_WEIGHTS = "EdgeVLM-Labs/mobile-videogpt-finetune-2000"
DEFAULT_VIDEOS_PER_EXERCISE = 20
DEFAULT_TEMPERATURES: List[Optional[float]] = [None, 0.2, 0.5, 0.7]  # None = greedy
PROMPT = "Please evaluate the exercise form shown. What mistakes, if any, are present, and what corrections would you recommend?"

# Metric thresholds
ROUGE_GREEN_THRESHOLD = 0.5
ROUGE_YELLOW_THRESHOLD = 0.2
BERT_GREEN_THRESHOLD = 0.7
BERT_YELLOW_THRESHOLD = 0.4
METEOR_GREEN_THRESHOLD = 0.5
METEOR_YELLOW_THRESHOLD = 0.2
LLM_GREEN_THRESHOLD = 4.0
LLM_YELLOW_THRESHOLD = 3.0

# LLM Judge prompts
LLM_JUDGE_SYSTEM_PROMPT = """You are an intelligent chatbot designed for evaluating feedback sequences provided by a virtual fitness coach to a person.

Your task is to compare the accuracy of the predicted feedback with the ground truth feedback.
- The predicted feedback must be factually accurate, relevant and align with the ground truth feedback.
- Consider synonyms or paraphrases as valid matches. Different wording that conveys the same corrective advice is equally correct.
- Repetition counts can be expressed in numeric form or in words (e.g. "10" and "ten" are equivalent).
- Encouraging, coaching-style language that addresses the same issue as the ground truth is valid.

## Scoring Rubric (1-5)

Score 5: Correctly identifies the exercise AND provides relevant, accurate corrective feedback that aligns with the ground truth. Paraphrases and synonyms are fully acceptable.
Score 4: Correctly identifies the exercise, feedback is relevant and mostly accurate but could be more specific or misses a minor detail.
Score 3: Exercise identified correctly, feedback is generic or vague but not factually wrong.
Score 2: Exercise identified but feedback is irrelevant or contradicts the ground truth.
Score 1: Wrong exercise identified or completely irrelevant response.

Provide your evaluation ONLY as a Python dictionary string with keys 'score' and 'justification'. Do not provide any other output text or explanation."""

LLM_JUDGE_USER_TEMPLATE = """Evaluate the following predicted feedback:
- Ground truth feedback: {ground_truth}
- Predicted feedback: {model_prediction}

Respond ONLY in this format: {{"score": <1-5>, "justification": "<one sentence>"}}"""


def temp_label(t: Optional[float]) -> str:
    """Human-readable label for a temperature setting."""
    if t is None:
        return "greedy"
    return f"temp_{t}"


# Model loading
def load_model(base_model: str, lora_path: str, device: str = "cuda"):
    """Load base model and merge LoRA adapters."""
    # Detect if lora_path contains LoRA adapters
    is_lora = False
    if "checkpoint-" in lora_path:
        is_lora = True
    elif os.path.isdir(lora_path) and os.path.exists(os.path.join(lora_path, "adapter_config.json")):
        is_lora = True
    else:
        try:
            hf_hub_download(lora_path, "adapter_config.json")
            is_lora = True
        except Exception:
            is_lora = False

    print(f"Base model : {base_model}")
    print(f"LoRA path  : {lora_path} (detected={'LoRA' if is_lora else 'full'})")

    config = AutoConfig.from_pretrained(base_model)
    tokenizer = AutoTokenizer.from_pretrained(base_model, use_fast=False)

    if is_lora:
        model = AutoModelForCausalLM.from_pretrained(
            base_model, config=config, torch_dtype=torch.float16
        )
        model = PeftModel.from_pretrained(model, lora_path)
        model = model.merge_and_unload()
    else:
        model = AutoModelForCausalLM.from_pretrained(
            lora_path, config=config, torch_dtype=torch.float16
        )

    model.to(device)
    return model, tokenizer


# Inference
def run_inference(
    model,
    tokenizer,
    video_path: str,
    prompt: str,
    device: str = "cuda",
    max_new_tokens: int = 512,
    temperature: Optional[float] = None,
) -> str:
    """Run inference on a single video at the given temperature."""
    input_ids, video_frames, context_frames, stop_str = preprocess_input(
        model, tokenizer, video_path, prompt
    )

    gen_kwargs = dict(
        images=torch.stack(video_frames, dim=0).half().to(device),
        context_images=torch.stack(context_frames, dim=0).half().to(device),
        num_beams=1,
        max_new_tokens=max_new_tokens,
        use_cache=True,
    )

    if temperature is None or temperature == 0.0:
        gen_kwargs["do_sample"] = False
    else:
        gen_kwargs["do_sample"] = True
        gen_kwargs["temperature"] = temperature

    with torch.inference_mode():
        output_ids = model.generate(input_ids, **gen_kwargs)

    output_text = tokenizer.batch_decode(output_ids, skip_special_tokens=True)[0].strip()
    if output_text.endswith(stop_str):
        output_text = output_text[: -len(stop_str)].strip()
    return output_text


# Dataset helpers
def load_ground_truths(labels_json: str, label_key: str = "labels_descriptive") -> Dict[str, str]:
    """Load fine_grained_labels.json and return {video_filename: ground_truth_text}."""

    with open(labels_json, "r") as f:
        data = json.load(f)

    mapping: Dict[str, str] = {}
    for entry in data:
        vp = entry["video_path"]  # e.g. './squats/00032607.mp4'
        filename = os.path.basename(vp)
        values = entry.get(label_key, entry.get("labels", []))
        mapping[filename] = "; ".join(values) if values else ""
    return mapping


def collect_videos(dataset_path: str, videos_per_exercise: int) -> List[Dict]:
    """Walk exercise subfolders and collect video paths with exercise type."""

    dataset_dir = Path(dataset_path)
    videos = []
    for subfolder in sorted(dataset_dir.iterdir()):
        if not subfolder.is_dir():
            continue
        exercise = subfolder.name
        mp4s = sorted(subfolder.glob("*.mp4"))[:videos_per_exercise]
        for mp4 in mp4s:
            videos.append(
                {
                    "video_path": str(mp4),
                    "exercise_type": exercise,
                    "filename": mp4.name,
                }
            )
    return videos


# Metric computation
def compute_meteor_score(reference: str, hypothesis: str, metric) -> float:
    if not reference or not hypothesis or metric is None:
        return 0.0
    try:
        return metric.compute(predictions=[hypothesis], references=[reference])["meteor"]
    except Exception:
        return 0.0


def compute_rouge_score(reference: str, hypothesis: str, metric) -> float:
    if not reference or not hypothesis or metric is None:
        return 0.0
    try:
        return metric.compute(predictions=[hypothesis], references=[reference])["rougeL"]
    except Exception:
        return 0.0


def compute_bert_similarity(text1: str, text2: str, model) -> float:
    if not text1 or not text2 or model is None:
        return 0.0
    try:
        from sklearn.metrics.pairwise import cosine_similarity as cos_sim

        embeddings = model.encode([text1, text2])
        return float(cos_sim([embeddings[0]], [embeddings[1]])[0][0])
    except Exception:
        return 0.0


def load_llm_judge():
    """Load Azure OpenAI GPT-4o as LLM judge."""
    try:
        from dotenv import load_dotenv
        from langchain_openai import AzureChatOpenAI

        load_dotenv(Path(__file__).parent.parent / ".env")
        llm = AzureChatOpenAI(
            model_name="gpt-4o",
            temperature=0.0,
            azure_endpoint=os.environ["AZURE_OPENAI_ENDPOINT"],
            azure_deployment=os.environ["AZURE_OPENAI_DEPLOYMENT_NAME"],
            api_key=os.environ["AZURE_OPENAI_API_KEY"],
            api_version=os.environ["OPENAI_API_VERSION"],
        )
        print("LLM Judge loaded (Azure OpenAI GPT-4o)")
        return llm
    except Exception as e:
        print(f"Warning: Could not load LLM judge: {e}")
        print("  LLM Accuracy scores will be skipped")
        return None


def compute_llm_accuracy_score(ground_truth: str, prediction: str, llm) -> float:
    if not ground_truth or not prediction or llm is None:
        return 0.0
    try:
        from langchain_core.messages import SystemMessage, HumanMessage

        user_prompt = LLM_JUDGE_USER_TEMPLATE.format(
            ground_truth=ground_truth,
            model_prediction=prediction,
        )
        response = llm.invoke(
            [
                SystemMessage(content=LLM_JUDGE_SYSTEM_PROMPT),
                HumanMessage(content=user_prompt),
            ]
        )
        response_text = response.content.strip()
        try:
            result = json.loads(response_text)
            score = float(result.get("score", 3))
            if 1 <= score <= 5:
                return score
        except (json.JSONDecodeError, TypeError, ValueError):
            pass
        match = re.search(r"\b([1-5])\b", response_text)
        if match:
            return float(match.group(1))
        return 3.0
    except Exception:
        return 0.0


# Excel report
def _color_fill(value: float, green_t: float, yellow_t: float):
    """Return an openpyxl PatternFill based on threshold."""
    from openpyxl.styles import PatternFill

    if value >= green_t:
        return PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    elif value >= yellow_t:
        return PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    else:
        return PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def write_excel(rows: List[Dict], temperatures: List[Optional[float]], output_path: str,
                measure_metrics: bool = False):
    """Write results to a formatted Excel workbook."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Temperature Sweep Results"

    # ---- Build header ----
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Fixed columns
    fixed_headers = ["Video Path", "Exercise Type", "Ground Truth"]
    # Per-temperature columns
    metric_names = ["BERT", "METEOR", "ROUGE-L", "LLM Judge"] if measure_metrics else []
    cols_per_temp = 1 + len(metric_names)  # prediction + metrics

    temp_headers = []
    for t in temperatures:
        label = temp_label(t)
        temp_headers.append(f"Pred ({label})")
        for m in metric_names:
            temp_headers.append(f"{m} ({label})")

    all_headers = fixed_headers + temp_headers

    for col_idx, header in enumerate(all_headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # ---- Write rows ----
    for row_idx, row_data in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=row_data["video_path"]).alignment = cell_align
        ws.cell(row=row_idx, column=2, value=row_data["exercise_type"]).alignment = cell_align
        ws.cell(row=row_idx, column=3, value=row_data["ground_truth"]).alignment = cell_align

        col = 4  # start after fixed columns
        for t in temperatures:
            label = temp_label(t)
            ws.cell(row=row_idx, column=col, value=row_data.get(f"pred_{label}", "")).alignment = cell_align
            col += 1

            if measure_metrics:
                # BERT
                val = row_data.get(f"bert_{label}", 0.0)
                c = ws.cell(row=row_idx, column=col, value=round(val, 4))
                c.fill = _color_fill(val, BERT_GREEN_THRESHOLD, BERT_YELLOW_THRESHOLD)
                c.alignment = cell_align
                col += 1

                # METEOR
                val = row_data.get(f"meteor_{label}", 0.0)
                c = ws.cell(row=row_idx, column=col, value=round(val, 4))
                c.fill = _color_fill(val, METEOR_GREEN_THRESHOLD, METEOR_YELLOW_THRESHOLD)
                c.alignment = cell_align
                col += 1

                # ROUGE-L
                val = row_data.get(f"rouge_{label}", 0.0)
                c = ws.cell(row=row_idx, column=col, value=round(val, 4))
                c.fill = _color_fill(val, ROUGE_GREEN_THRESHOLD, ROUGE_YELLOW_THRESHOLD)
                c.alignment = cell_align
                col += 1

                # LLM Judge
                val = row_data.get(f"llm_{label}", 0.0)
                c = ws.cell(row=row_idx, column=col, value=round(val, 2))
                c.fill = _color_fill(val, LLM_GREEN_THRESHOLD, LLM_YELLOW_THRESHOLD)
                c.alignment = cell_align
                col += 1

        # Apply border to all cells in row
        for c_idx in range(1, col):
            ws.cell(row=row_idx, column=c_idx).border = thin_border

    if measure_metrics:
        summary_row = len(rows) + 3
        ws.cell(row=summary_row, column=1, value="AVERAGES").font = Font(bold=True, size=12)

        col = 4
        for t in temperatures:
            label = temp_label(t)
            col += 1  # skip prediction column

            for metric_prefix in ["bert", "meteor", "rouge", "llm"]:
                key = f"{metric_prefix}_{label}"
                vals = [r.get(key, 0.0) for r in rows if r.get(key, 0.0) > 0]
                avg = np.mean(vals) if vals else 0.0
                c = ws.cell(row=summary_row, column=col, value=round(avg, 4))
                c.font = Font(bold=True)
                col += 1

    # ---- Column widths ----
    ws.column_dimensions[get_column_letter(1)].width = 40  # video path
    ws.column_dimensions[get_column_letter(2)].width = 30  # exercise type
    ws.column_dimensions[get_column_letter(3)].width = 50  # ground truth
    for c_idx in range(4, len(all_headers) + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 18

    # Widen prediction columns
    col = 4
    for _ in temperatures:
        ws.column_dimensions[get_column_letter(col)].width = 50
        col += cols_per_temp

    # Freeze header row and fixed columns
    ws.freeze_panes = "D2"

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"\nReport saved to: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Temperature sweep inference experiment")
    parser.add_argument("--dataset_path", type=str, default=DEFAULT_DATASET,
                        help="Path to dataset with exercise subfolders")
    parser.add_argument("--base_model", type=str, default=DEFAULT_BASE_MODEL)
    parser.add_argument("--lora_weights", type=str, default=DEFAULT_LORA_WEIGHTS)
    parser.add_argument("--videos_per_exercise", type=int, default=DEFAULT_VIDEOS_PER_EXERCISE)
    parser.add_argument("--output", type=str, default=None,
                        help="Output xlsx path (auto-generated from lora_weights if not provided)")
    parser.add_argument("--device", type=str, default="cuda")
    parser.add_argument("--max_new_tokens", type=int, default=512)
    parser.add_argument("--label_key", type=str, default="labels_descriptive",
                        help="Key to extract ground truth from fine_grained_labels.json "
                             "(e.g. 'labels_descriptive', 'labels')")
    parser.add_argument("--measure_metrics", action="store_true",
                        help="Compute evaluation metrics (BERT, METEOR, ROUGE-L, LLM Judge). "
                             "Off by default; when disabled only predictions are generated.")
    parser.add_argument("--no_llm_judge", action="store_true", help="Skip LLM judge metric (only relevant with --measure_metrics)")
    args = parser.parse_args()

    temperatures = DEFAULT_TEMPERATURES

    if args.output is None:
        # Extract repo name from HF link (e.g. "EdgeVLM-Labs/mobile-videogpt-finetune-2000" -> "mobile-videogpt-finetune-2000")
        lora_name = args.lora_weights.rstrip("/").split("/")[-1]
        args.output = f"results/vary-inference/vary_inference_report-{lora_name}.xlsx"

    labels_json = os.path.join(args.dataset_path, "fine_grained_labels.json")
    if not os.path.exists(labels_json):
        print(f"Error: {labels_json} not found")
        sys.exit(1)
    gt_map = load_ground_truths(labels_json, label_key=args.label_key)
    print(f"Loaded {len(gt_map)} ground truth entries (key='{args.label_key}')")

    videos = collect_videos(args.dataset_path, args.videos_per_exercise)
    print(f"Collected {len(videos)} videos across "
          f"{len(set(v['exercise_type'] for v in videos))} exercises "
          f"(max {args.videos_per_exercise} per exercise)")

    print("\nLoading model...")
    model, tokenizer = load_model(args.base_model, args.lora_weights, args.device)

    meteor_metric = None
    rouge_metric = None
    bert_model = None
    llm_judge = None

    if args.measure_metrics:
        import evaluate

        print("\nLoading evaluation metrics...")
        meteor_metric = evaluate.load("meteor")
        rouge_metric = evaluate.load("rouge")

        try:
            from sentence_transformers import SentenceTransformer
            bert_model = SentenceTransformer("all-MiniLM-L6-v2")
            print("BERT model loaded")
        except Exception as e:
            print(f"Warning: Could not load BERT model: {e}")

        if not args.no_llm_judge:
            llm_judge = load_llm_judge()
    else:
        print("\nMetrics disabled (use --measure_metrics to enable)")

    results: List[Dict] = []
    total_inferences = len(videos) * len(temperatures)
    print(f"\nRunning {total_inferences} inferences "
          f"({len(videos)} videos x {len(temperatures)} temperatures)...\n")

    for vid_info in tqdm(videos, desc="Videos"):
        video_path = vid_info["video_path"]
        exercise = vid_info["exercise_type"]
        filename = vid_info["filename"]
        ground_truth = gt_map.get(filename, "")

        row: Dict = {
            "video_path": video_path,
            "exercise_type": exercise,
            "ground_truth": ground_truth,
        }

        for t in temperatures:
            label = temp_label(t)
            try:
                prediction = run_inference(
                    model, tokenizer, video_path, PROMPT,
                    args.device, args.max_new_tokens, temperature=t,
                )
            except Exception as e:
                print(f"\n  Error: {filename} @ {label}: {e}")
                prediction = ""

            row[f"pred_{label}"] = prediction

            # Compute metrics against ground truth (only if enabled)
            if args.measure_metrics:
                row[f"bert_{label}"] = compute_bert_similarity(ground_truth, prediction, bert_model)
                row[f"meteor_{label}"] = compute_meteor_score(ground_truth, prediction, meteor_metric)
                row[f"rouge_{label}"] = compute_rouge_score(ground_truth, prediction, rouge_metric)
                row[f"llm_{label}"] = compute_llm_accuracy_score(ground_truth, prediction, llm_judge)

        results.append(row)

    write_excel(results, temperatures, args.output, measure_metrics=args.measure_metrics)

    print(f"\n{'='*60}")
    print("Temperature Sweep Complete")
    print(f"{'='*60}")
    print(f"Videos processed : {len(results)}")
    print(f"Temperatures     : {[temp_label(t) for t in temperatures]}")
    print(f"Metrics          : {'enabled' if args.measure_metrics else 'disabled'}")

    if args.measure_metrics:
        for t in temperatures:
            label = temp_label(t)
            bert_vals = [r[f"bert_{label}"] for r in results if r.get(f"bert_{label}", 0) > 0]
            meteor_vals = [r[f"meteor_{label}"] for r in results if r.get(f"meteor_{label}", 0) > 0]
            rouge_vals = [r[f"rouge_{label}"] for r in results if r.get(f"rouge_{label}", 0) > 0]
            llm_vals = [r[f"llm_{label}"] for r in results if r.get(f"llm_{label}", 0) > 0]

            print(f"\n  [{label}]")
            if bert_vals:
                print(f"    BERT   mean={np.mean(bert_vals):.4f}")
            if meteor_vals:
                print(f"    METEOR mean={np.mean(meteor_vals):.4f}")
            if rouge_vals:
                print(f"    ROUGE  mean={np.mean(rouge_vals):.4f}")
            if llm_vals:
                print(f"    LLM    mean={np.mean(llm_vals):.2f}")

    print(f"\nReport: {args.output}")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
